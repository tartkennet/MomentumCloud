#!/usr/bin/env python3
from __future__ import annotations

import json
import logging
import math
import os
import random
import re
import time
from pathlib import Path
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

POLYGON_BASE = "https://api.polygon.io"
INPUT_TEMPLATE = "Momentum mall fil.xlsx"
OUTPUT_FILE = "Momentum.xlsx"
UNIVERSE_SHEET = "Universe"
PARAM_SHEET = "Parametrar"
BUYNOW_SHEET = "BuyNow"
CANDIDATES_SHEET = "Kandidater"
DIAGNOSTICS_SHEET = "Diagnostics"
DEBUG_SHEET = "DebugTickers"
TRACKING_SHEET = "Utvärdering"
TRACKING_FILE = "Momentum_buy_history.csv"
PORTFOLIO_FILE = "Momentum_portfolio.csv"
HOLDINGS_SHEET = "Holdings"
PORTFOLIO_SUMMARY_SHEET = "Portföljöversikt"
HISTORY_CACHE_FILE = ".momentum_history_cache.json"
SPLITS_CACHE_FILE = ".momentum_splits_cache.json"
HISTORY_CACHE_VERSION = 2

MIN_PRICE = 5.0
MIN_ADTV20 = 20_000_000.0
TOP_N = 15
TOP10_N = 10


def safe_error_text(exc: Exception) -> str:
    """Return a useful error message without leaking Polygon credentials."""
    return re.sub(
        r"(?i)([?&]apiKey=)[^&\s]+",
        r"\1[REDACTED]",
        str(exc),
    )


@dataclass
class StockMetrics:
    ticker: str
    price_usd: float
    avg_dollar_volume20: float
    sma50: float
    sma200: float
    ret5d: float
    ret1m: float
    ret3m: float
    ret6m: float
    ret12m: float
    accel: float
    dist_sma50: float
    total_score: float


class PolygonClient:
    def __init__(self, api_key: str) -> None:
        self.api_key = api_key
        self.session = requests.Session()
        self.grouped_cache: Dict[str, Dict[str, Dict[str, float]]] = {}

    def _get(self, path: str, params: Optional[Dict] = None) -> Dict:
        merged = {"apiKey": self.api_key}
        if params:
            merged.update(params)
        url = f"{POLYGON_BASE}{path}"
        max_attempts = int(os.getenv("POLYGON_MAX_RETRIES", "40"))
        rate_limit_wait = int(os.getenv("POLYGON_RATE_LIMIT_WAIT_SECONDS", "60"))
        for attempt in range(1, max_attempts + 1):
            try:
                response = self.session.get(url, params=merged, timeout=30)
                if response.status_code == 429:
                    retry_after = response.headers.get("Retry-After")
                    try:
                        wait = max(rate_limit_wait, int(float(retry_after))) if retry_after else rate_limit_wait
                    except ValueError:
                        wait = rate_limit_wait
                    # Add jitter so simultaneous workflows do not retry in lockstep
                    # and keep starving the same request.
                    wait = float(wait) + random.uniform(0.0, max(1.0, float(wait) * 0.25))
                    logging.warning(
                        "Polygon rate limit. Waiting %.1fs and retrying (%s/%s).",
                        wait, attempt, max_attempts,
                    )
                    time.sleep(wait)
                    continue
                response.raise_for_status()
                data = response.json()
                if data.get("status") == "ERROR":
                    raise RuntimeError(f"Polygon error: {data}")
                return data
            except requests.RequestException:
                if attempt >= max_attempts:
                    raise
                wait = min(30, 3 * attempt)
                logging.warning(
                    "Polygon request failed. Waiting %ss and retrying (%s/%s).",
                    wait, attempt, max_attempts,
                )
                time.sleep(wait)
        raise RuntimeError("Polygon request failed after retry budget")

    def fetch_fx_close(self, date_str: str) -> float:
        data = self._get(
            f"/v2/aggs/ticker/C:USDSEK/range/1/day/{date_str}/{date_str}",
            {"adjusted": "true", "limit": 10},
        )
        results = data.get("results", []) or []
        if not results:
            raise ValueError(f"No USDSEK result for {date_str}")
        return float(results[-1]["c"])

    def fetch_recent_splits(self, start_date: str, end_date: str) -> List[Dict]:
        data = self._get(
            "/stocks/v1/splits",
            {"execution_date.gte": start_date, "execution_date.lte": end_date, "limit": 5000, "sort": "execution_date.asc"},
        )
        return data.get("results", []) or []

    def fetch_grouped_stocks(self, date_str: str) -> Dict[str, Dict[str, float]]:
        if date_str in self.grouped_cache:
            return self.grouped_cache[date_str]
        data = self._get(
            f"/v2/aggs/grouped/locale/us/market/stocks/{date_str}",
            {"adjusted": "true"},
        )
        out: Dict[str, Dict[str, float]] = {}
        for row in data.get("results", []) or []:
            ticker = str(row.get("T", "")).upper().strip()
            close = row.get("c")
            volume = row.get("v")
            if ticker and close is not None and volume is not None:
                out[ticker] = {
                    "open": float(row["o"]) if row.get("o") is not None else float(close),
                    "high": float(row["h"]) if row.get("h") is not None else float(close),
                    "low": float(row["l"]) if row.get("l") is not None else float(close),
                    "close": float(close),
                    "volume": float(volume),
                }
        self.grouped_cache[date_str] = out
        return out


def setup_logging() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")


def normalize_ticker(ticker: str) -> str:
    return str(ticker).strip().upper()


def alt_ticker(ticker: str) -> str:
    return normalize_ticker(ticker).replace(".", "-")


def get_from_map(day_map: Dict[str, Dict[str, float]], ticker: str) -> Optional[Dict[str, float]]:
    return day_map.get(normalize_ticker(ticker)) or day_map.get(alt_ticker(ticker))


def read_universe(path: str) -> List[str]:
    df = pd.read_excel(path, sheet_name=UNIVERSE_SHEET)
    if "Ticker" not in df.columns:
        raise ValueError("Universe sheet must contain a 'Ticker' column")
    tickers = df["Ticker"].dropna().astype(str).str.strip().str.upper().replace("", np.nan).dropna().tolist()
    unique = list(dict.fromkeys(tickers))
    if not unique:
        raise ValueError("No tickers found in Universe sheet")
    return unique


def read_portfolio_value_sek(path: str) -> float:
    wb = load_workbook(path, data_only=True)
    value = wb[PARAM_SHEET]["B1"].value
    if value is None:
        raise ValueError("Parametrar!B1 is empty")
    parsed = float(value)
    if parsed <= 0:
        raise ValueError("Parametrar!B1 must be > 0")
    return parsed


def pick_latest_grouped_date(client: PolygonClient, max_lookback_days: int = 10) -> Tuple[str, Dict[str, Dict[str, float]], float]:
    today = datetime.now(timezone.utc).date()
    # The production workflow runs after the US close, so try the current UTC
    # date first. Polygon returns no complete grouped result before it exists.
    for days_back in range(0, max_lookback_days + 1):
        date_str = (today - timedelta(days=days_back)).isoformat()
        try:
            grouped = client.fetch_grouped_stocks(date_str)
            fx = client.fetch_fx_close(date_str)
            if len(grouped) > 1000:
                return date_str, grouped, fx
        except Exception as exc:
            logging.info(
                "Grouped date %s unavailable: %s",
                date_str,
                safe_error_text(exc),
            )
    raise RuntimeError("Could not find recent grouped stock data with USDSEK")


def slim_grouped_map(grouped: Dict[str, Dict[str, float]], tickers: List[str]) -> Dict[str, Dict[str, float]]:
    relevant = {normalize_ticker(ticker) for ticker in tickers}
    relevant.update(alt_ticker(ticker) for ticker in tickers)
    return {ticker: row for ticker, row in grouped.items() if ticker in relevant}


def load_history_cache_state(
    path: str,
) -> Tuple[Dict[str, Dict[str, Dict[str, float]]], Optional[str]]:
    """Load only history caches whose split-adjustment basis is known."""
    cache_path = Path(path)
    if not cache_path.exists():
        return {}, None
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
        sessions = payload.get("sessions")
        adjusted_through = payload.get("adjusted_through")
        if (
            payload.get("version") != HISTORY_CACHE_VERSION
            or not isinstance(sessions, dict)
            or not adjusted_through
        ):
            raise ValueError("legacy cache has an unknown split-adjustment basis")
        return sessions, str(adjusted_through)
    except Exception as exc:
        logging.warning("Discarding unsafe Momentum history cache: %s", exc)
        return {}, None


def load_history_cache(path: str) -> Dict[str, Dict[str, Dict[str, float]]]:
    """Backward-compatible sessions-only cache reader."""
    sessions, _ = load_history_cache_state(path)
    return sessions


def save_history_cache(
    path: str,
    maps: List[Tuple[str, Dict[str, Dict[str, float]]]],
    adjusted_through: Optional[str] = None,
) -> None:
    basis_date = adjusted_through or datetime.now(timezone.utc).date().isoformat()
    payload = {
        "version": HISTORY_CACHE_VERSION,
        "adjusted_through": basis_date,
        "sessions": {date_str: grouped for date_str, grouped in maps},
    }
    Path(path).write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")


def adjust_history_cache_for_splits(
    sessions: Dict[str, Dict[str, Dict[str, float]]],
    split_events: List[Dict],
    adjusted_through: Optional[str],
    target_basis_date: str,
) -> int:
    """Restate cached OHLCV onto the latest split-adjusted price basis."""
    if not sessions or not adjusted_through:
        return 0

    applied = 0
    for event in sorted(split_events, key=lambda item: str(item.get("execution_date", ""))):
        event_day = str(event.get("execution_date", ""))
        if not (str(adjusted_through) < event_day <= str(target_basis_date)):
            continue
        split_from = float(event.get("split_from", 0) or 0)
        split_to = float(event.get("split_to", 0) or 0)
        ticker = normalize_ticker(str(event.get("ticker", "")).upper().strip())
        if not ticker or split_from <= 0 or split_to <= 0:
            continue

        price_factor = split_from / split_to
        volume_factor = split_to / split_from
        touched = 0
        for session_day, day_map in sessions.items():
            if session_day >= event_day:
                continue
            row = get_from_map(day_map, ticker)
            if not row:
                continue
            for field in ("open", "high", "low", "close"):
                if field in row and row[field] is not None:
                    row[field] = float(row[field]) * price_factor
            if "volume" in row and row["volume"] is not None:
                row["volume"] = float(row["volume"]) * volume_factor
            touched += 1
        if touched:
            logging.info(
                "Adjusted %d cached %s rows for split %s:%s on %s",
                touched,
                ticker,
                split_from,
                split_to,
                event_day,
            )
            applied += 1
    return applied


def collect_recent_trading_maps(
    client: PolygonClient,
    latest_date: str,
    tickers: List[str],
    latest_map: Dict[str, Dict[str, float]],
    needed: int = 253,
    cache_file: str = HISTORY_CACHE_FILE,
) -> List[Tuple[str, Dict[str, Dict[str, float]]]]:
    latest = datetime.fromisoformat(latest_date).date()
    cached_maps, adjusted_through = load_history_cache_state(cache_file)
    cache_basis_date = datetime.now(timezone.utc).date().isoformat()
    if cached_maps and adjusted_through and adjusted_through < cache_basis_date:
        new_splits = client.fetch_recent_splits(adjusted_through, cache_basis_date)
        adjust_history_cache_for_splits(
            cached_maps,
            new_splits,
            adjusted_through,
            cache_basis_date,
        )
    cached_maps[latest_date] = slim_grouped_map(latest_map, tickers)
    maps: List[Tuple[str, Dict[str, Dict[str, float]]]] = []
    days_back = 0
    max_calendar_days = int(needed * 2.5) + 30
    while len(maps) < needed and days_back < max_calendar_days:
        candidate_date = latest - timedelta(days=days_back)
        days_back += 1
        if candidate_date.weekday() >= 5:
            continue
        date_str = candidate_date.isoformat()
        cached = cached_maps.get(date_str)
        if cached is not None:
            maps.append((date_str, cached))
            continue
        try:
            grouped = client.fetch_grouped_stocks(date_str)
            if len(grouped) > 1000:
                cached_maps[date_str] = slim_grouped_map(grouped, tickers)
                maps.append((date_str, cached_maps[date_str]))
        except Exception:
            pass
    if len(maps) < needed:
        raise RuntimeError(f"Only found {len(maps)} recent trading maps, need {needed}")
    save_history_cache(cache_file, maps, cache_basis_date)
    return maps


def pct_return(current: float, old: Optional[float]) -> Optional[float]:
    if old is None or old <= 0:
        return None
    return (current / old - 1.0) * 100.0


def simple_moving_average(closes: List[float], period: int) -> Optional[float]:
    if len(closes) < period:
        return None
    return float(np.mean(closes[:period]))


def build_metrics(tickers, latest_map, recent253):
    counters = {
        "total_tickers_in_universe": len(tickers),
        "grouped_data_pipeline": 1,
        "latest_grouped_symbols": len(latest_map),
        "successful_polygon_data": 0,
        "no_data_from_polygon": 0,
        "excluded_by_insufficient_history": 0,
        "excluded_by_price_filter": 0,
        "excluded_by_liquidity_filter": 0,
        "excluded_by_trend_filter": 0,
        "excluded_by_momentum_filter": 0,
        "excluded_by_negative_5d_filter": 0,
        "excluded_by_decelerating_momentum_filter": 0,
        "final_candidate_count": 0,
    }
    metrics: List[StockMetrics] = []
    debug_rows = []

    for ticker in tickers:
        latest = get_from_map(latest_map, ticker)
        if not latest:
            counters["no_data_from_polygon"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "NO_LATEST_DATA"})
            continue

        counters["successful_polygon_data"] += 1
        price = float(latest["close"])
        if price < MIN_PRICE:
            counters["excluded_by_price_filter"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "PRICE", "LastClose": price})
            continue

        dollar_values = []
        for _, grouped in recent253[:20]:
            row = get_from_map(grouped, ticker)
            if row:
                dollar_values.append(float(row["close"]) * float(row["volume"]))
        if len(dollar_values) < 10:
            counters["excluded_by_insufficient_history"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "INSUFFICIENT_20D", "RowsReturned": len(dollar_values), "LastClose": price})
            continue

        adv20 = float(np.mean(dollar_values))
        if adv20 < MIN_ADTV20:
            counters["excluded_by_liquidity_filter"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "LIQUIDITY", "LastClose": price, "AvgDollarVolume20": adv20})
            continue

        closes = []
        for _, grouped in recent253:
            row = get_from_map(grouped, ticker)
            if not row:
                break
            closes.append(float(row["close"]))

        if len(closes) < 253:
            counters["excluded_by_insufficient_history"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "INSUFFICIENT_253D", "LastClose": price, "RowsReturned": len(closes)})
            continue

        c5 = closes[5]
        c21 = closes[21]
        c63 = closes[63]
        c126 = closes[126]
        c252 = closes[252]
        sma50 = simple_moving_average(closes, 50)
        sma200 = simple_moving_average(closes, 200)

        if sma50 is None or sma200 is None:
            counters["excluded_by_insufficient_history"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "MISSING_SMA_HISTORY", "LastClose": price})
            continue
        if not (price > sma50 and price > sma200):
            counters["excluded_by_trend_filter"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "TREND", "LastClose": price, "SMA50": sma50, "SMA200": sma200})
            continue

        ret5d = pct_return(price, c5)
        ret1m = pct_return(price, c21)
        ret3m = pct_return(price, c63)
        ret6m = pct_return(price, c126)
        ret12m = pct_return(price, c252)
        if None in [ret5d, ret1m, ret3m, ret6m, ret12m]:
            counters["excluded_by_insufficient_history"] += 1
            continue

        if ret5d <= -3:
            counters["excluded_by_negative_5d_filter"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "NEGATIVE_5D", "LastClose": price, "Ret5d": ret5d})
            continue

        if not (ret1m > 0 and ret3m > 0 and ret6m > 0 and ret12m > 0):
            counters["excluded_by_momentum_filter"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "MOMENTUM", "LastClose": price, "Ret1m": ret1m, "Ret3m": ret3m, "Ret6m": ret6m, "Ret12m": ret12m})
            continue

        dist_sma50 = (price / sma50 - 1.0) * 100.0

        if ret1m <= ret3m / 3.0:
            counters["excluded_by_decelerating_momentum_filter"] += 1
            debug_rows.append({"Ticker": ticker, "Status": "DECELERATING", "LastClose": price, "Ret1m": ret1m, "Ret3m": ret3m, "Dist_SMA50": dist_sma50})
            continue

        accel = float(ret1m - ret3m)
        total_score = (
            0.35 * ret1m
            + 0.30 * ret3m
            + 0.20 * ret6m
            + 0.10 * ret12m
            + 0.15 * accel
            - 0.50 * max(dist_sma50 - 8.0, 0.0)
        )

        metrics.append(StockMetrics(ticker, price, adv20, sma50, sma200, ret5d, ret1m, ret3m, ret6m, ret12m, accel, dist_sma50, total_score))
        debug_rows.append({"Ticker": ticker, "Status": "PASS", "LastClose": price, "AvgDollarVolume20": adv20, "Ret5d": ret5d, "Ret1m": ret1m, "Ret3m": ret3m, "Ret6m": ret6m, "Ret12m": ret12m, "Dist_SMA50": dist_sma50, "Accel": accel, "TotalScore": total_score})

    counters["final_candidate_count"] = len(metrics)
    return metrics, counters, pd.DataFrame(debug_rows[:300])



def load_split_cache(path: str = SPLITS_CACHE_FILE) -> Dict:
    cache_path = Path(path)
    if not cache_path.exists():
        return {"last_checked": None, "events": {}}
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
        if not isinstance(payload.get("events"), dict):
            raise ValueError("invalid split cache")
        return payload
    except Exception:
        return {"last_checked": None, "events": {}}


def update_split_cache(client: PolygonClient, latest_date: str, path: str = SPLITS_CACHE_FILE) -> Dict[str, List[Dict]]:
    payload = load_split_cache(path)
    previous = payload.get("last_checked")
    start_date = previous or (datetime.fromisoformat(latest_date).date() - timedelta(days=14)).isoformat()
    events = payload.get("events", {})
    for event in client.fetch_recent_splits(start_date, latest_date):
        ticker = normalize_ticker(event.get("ticker", ""))
        event_id = str(event.get("id", ""))
        if not ticker or not event_id:
            continue
        known = {str(item.get("id", "")) for item in events.get(ticker, [])}
        if event_id not in known:
            events.setdefault(ticker, []).append(event)
    Path(path).write_text(json.dumps({"last_checked": latest_date, "events": events}, separators=(",", ":")), encoding="utf-8")
    return events


def split_price_factor(events: List[Dict], entry_date: str, latest_date: str) -> float:
    factor = 1.0
    for event in events:
        execution_date = str(event.get("execution_date", ""))
        if entry_date < execution_date <= latest_date:
            split_from = float(event.get("split_from", 0) or 0)
            split_to = float(event.get("split_to", 0) or 0)
            if split_from > 0 and split_to > 0:
                factor *= split_from / split_to
    return factor

def update_buy_history(
    buynow_df: pd.DataFrame,
    latest_map: Dict[str, Dict[str, float]],
    latest_date: str,
    tracking_file: str = TRACKING_FILE,
    split_events_by_ticker: Optional[Dict[str, List[Dict]]] = None,
) -> pd.DataFrame:
    columns = [
        "EntryDate", "Ticker", "EntryRank", "EntryPrice_USD", "EntryTotalScore",
        "LatestDate", "LatestPrice_USD", "AdjustedEntryPrice_USD", "Return_Since_Entry_pct",
        "Price_1W", "Return_1W_pct", "Price_2W", "Return_2W_pct", "Price_3W", "Return_3W_pct",
    ]
    if os.path.exists(tracking_file):
        history = pd.read_csv(tracking_file)
    else:
        history = pd.DataFrame(columns=columns)
    history = history.reindex(columns=columns)

    for row_idx, row in history.iterrows():
        latest = get_from_map(latest_map, str(row["Ticker"]))
        if not latest:
            continue
        price = float(latest["close"])
        history.at[row_idx, "LatestDate"] = latest_date
        history.at[row_idx, "LatestPrice_USD"] = price
        entry_price = float(row["EntryPrice_USD"])
        events = (split_events_by_ticker or {}).get(normalize_ticker(str(row["Ticker"])), [])
        adjusted_entry_price = entry_price * split_price_factor(events, str(row["EntryDate"]), latest_date)
        history.at[row_idx, "AdjustedEntryPrice_USD"] = adjusted_entry_price
        history.at[row_idx, "Return_Since_Entry_pct"] = pct_return(price, adjusted_entry_price)
        entry_date = pd.Timestamp(row["EntryDate"])
        for days, price_col, return_col in ((7, "Price_1W", "Return_1W_pct"), (14, "Price_2W", "Return_2W_pct"), (21, "Price_3W", "Return_3W_pct")):
            if pd.Timestamp(latest_date) >= entry_date + pd.Timedelta(days=days) and pd.isna(row[price_col]):
                history.at[row_idx, price_col] = price
                history.at[row_idx, return_col] = pct_return(price, entry_price)

    existing = set(zip(history["EntryDate"].astype(str), history["Ticker"].astype(str)))
    new_rows = []
    for _, row in buynow_df.iterrows():
        key = (latest_date, str(row["Ticker"]))
        if key not in existing:
            price = float(row["Price_USD"])
            new_rows.append({
                "EntryDate": latest_date, "Ticker": row["Ticker"], "EntryRank": int(row["Rank"]),
                "EntryPrice_USD": price, "EntryTotalScore": float(row["TotalScore"]),
                "LatestDate": latest_date, "LatestPrice_USD": price, "Return_Since_Entry_pct": 0.0,
            })
    if new_rows:
        history = pd.concat([history, pd.DataFrame(new_rows)], ignore_index=True)
    history.to_csv(tracking_file, index=False)
    return history.sort_values(["EntryDate", "EntryRank"], ascending=[False, True], na_position="last").reset_index(drop=True)



def load_portfolio(path: str = PORTFOLIO_FILE) -> pd.DataFrame:
    columns = ["Ticker", "Qty", "EntryDate", "EntryPrice_USD", "EntryFX", "EntryFees_SEK"]
    if not os.path.exists(path):
        return pd.DataFrame(columns=columns)
    portfolio = pd.read_csv(path)
    missing = [column for column in columns if column not in portfolio.columns]
    if missing:
        raise ValueError(f"{path} is missing columns: {', '.join(missing)}")
    portfolio = portfolio.reindex(columns=columns).dropna(subset=["Ticker", "Qty", "EntryPrice_USD", "EntryFX"])
    portfolio["Ticker"] = portfolio["Ticker"].map(normalize_ticker)
    for column in ("Qty", "EntryPrice_USD", "EntryFX", "EntryFees_SEK"):
        portfolio[column] = pd.to_numeric(portfolio[column], errors="coerce")
    portfolio["EntryFees_SEK"] = portfolio["EntryFees_SEK"].fillna(0.0)
    return portfolio.dropna(subset=["Qty", "EntryPrice_USD", "EntryFX"])


def close_on_or_before(recent_maps, ticker: str, date_value) -> Optional[float]:
    try:
        target = pd.Timestamp(date_value).date()
    except Exception:
        return None
    eligible = [
        (date_str, get_from_map(day_map, ticker))
        for date_str, day_map in recent_maps
        if datetime.fromisoformat(date_str).date() <= target
    ]
    if not eligible:
        return None
    date_str, row = max(eligible, key=lambda item: item[0])
    return float(row["close"]) if row else None


def build_holdings_df(portfolio_df, latest_map, fx: float, candidates_df, recent_maps) -> pd.DataFrame:
    columns = [
        "Ticker", "Qty", "EntryDate", "EntryPrice_USD", "EntryFX", "EntryFees_SEK",
        "EntryValue_SEK", "Price_USD", "FX", "Value_SEK", "P&L_SEK", "P&L_pct",
        "SPY_Return_SEK_pct", "Edge_vs_SPY_pct", "CurrentRank", "MomentumStatus",
    ]
    if portfolio_df.empty:
        return pd.DataFrame(columns=columns)

    rank_lookup = candidates_df.set_index("Ticker")["Rank"].to_dict() if not candidates_df.empty else {}
    current_spy = get_from_map(latest_map, "SPY")
    rows = []
    for _, holding in portfolio_df.iterrows():
        ticker = normalize_ticker(holding["Ticker"])
        latest = get_from_map(latest_map, ticker)
        entry_value = float(holding["Qty"]) * float(holding["EntryPrice_USD"]) * float(holding["EntryFX"]) + float(holding["EntryFees_SEK"])
        row = {
            "Ticker": ticker,
            "Qty": float(holding["Qty"]),
            "EntryDate": holding["EntryDate"],
            "EntryPrice_USD": float(holding["EntryPrice_USD"]),
            "EntryFX": float(holding["EntryFX"]),
            "EntryFees_SEK": float(holding["EntryFees_SEK"]),
            "EntryValue_SEK": entry_value,
            "Price_USD": None,
            "FX": fx,
            "Value_SEK": None,
            "P&L_SEK": None,
            "P&L_pct": None,
            "SPY_Return_SEK_pct": None,
            "Edge_vs_SPY_pct": None,
            "CurrentRank": rank_lookup.get(ticker),
            "MomentumStatus": "NO_PRICE",
        }
        if latest:
            price = float(latest["close"])
            value = float(holding["Qty"]) * price * fx
            pnl = value - entry_value
            row.update({"Price_USD": price, "Value_SEK": value, "P&L_SEK": pnl, "P&L_pct": pct_return(value, entry_value)})
            rank = rank_lookup.get(ticker)
            row["MomentumStatus"] = "TOP10" if rank and rank <= TOP10_N else ("BUY_NOW" if rank and rank <= TOP_N else ("CANDIDATE" if rank else "OUT_OF_MOMENTUM"))
            entry_spy = close_on_or_before(recent_maps, "SPY", holding["EntryDate"])
            if current_spy and entry_spy:
                spy_return = pct_return(float(current_spy["close"]) * fx, entry_spy * float(holding["EntryFX"]))
                row["SPY_Return_SEK_pct"] = spy_return
                if row["P&L_pct"] is not None and spy_return is not None:
                    row["Edge_vs_SPY_pct"] = row["P&L_pct"] - spy_return
        rows.append(row)
    return pd.DataFrame(rows, columns=columns)


def build_portfolio_summary(holdings_df, latest_date: str) -> pd.DataFrame:
    if holdings_df.empty:
        return pd.DataFrame([{"Metric": "Status", "Value": "No holdings imported yet"}])
    entry_value = float(holdings_df["EntryValue_SEK"].sum())
    current_value = float(holdings_df["Value_SEK"].sum(min_count=1))
    pnl = current_value - entry_value if not pd.isna(current_value) else None
    return pd.DataFrame([
        {"Metric": "Latest market date", "Value": latest_date},
        {"Metric": "Positions", "Value": int(len(holdings_df))},
        {"Metric": "Entry value SEK", "Value": entry_value},
        {"Metric": "Current value SEK", "Value": current_value},
        {"Metric": "P&L SEK", "Value": pnl},
        {"Metric": "P&L %", "Value": pct_return(current_value, entry_value) if pnl is not None else None},
        {"Metric": "Benchmark", "Value": "SPY, SEK adjusted from each entry date"},
    ])

def metrics_to_ranked_df(metrics: List[StockMetrics], fx: float, portfolio_value_sek: float, limit: Optional[int] = None) -> pd.DataFrame:
    ranked = sorted(metrics, key=lambda x: x.total_score, reverse=True)
    if limit is not None:
        ranked = ranked[:limit]

    rows = []
    position_size_sek = portfolio_value_sek / TOP10_N
    for rank, metric in enumerate(ranked, start=1):
        is_top10 = rank <= TOP10_N
        price_sek = metric.price_usd * fx
        qty = max(1, math.floor(position_size_sek / price_sek)) if is_top10 else ""
        rows.append({
            "Rank": rank,
            "Ticker": metric.ticker,
            "Top10": bool(is_top10),
            "Qty": qty,
            "Price_USD": metric.price_usd,
            "FX": fx,
            "Price_SEK": price_sek,
            "Position_SEK": qty * price_sek if is_top10 else "",
            "AvgDollarVolume20": metric.avg_dollar_volume20,
            "SMA50": metric.sma50,
            "SMA200": metric.sma200,
            "Ret5d": metric.ret5d,
            "Ret1m": metric.ret1m,
            "Ret3m": metric.ret3m,
            "Ret6m": metric.ret6m,
            "Ret12m": metric.ret12m,
            "Dist_SMA50": metric.dist_sma50,
            "Accel": metric.accel,
            "TotalScore": metric.total_score,
        })
    return pd.DataFrame(rows)


def build_buynow_df(metrics: List[StockMetrics], fx: float, portfolio_value_sek: float) -> pd.DataFrame:
    return metrics_to_ranked_df(metrics, fx, portfolio_value_sek, limit=TOP_N)


def build_candidates_df(metrics: List[StockMetrics], fx: float, portfolio_value_sek: float) -> pd.DataFrame:
    return metrics_to_ranked_df(metrics, fx, portfolio_value_sek, limit=None)


def autosize_columns(ws) -> None:
    if ws.max_row < 1:
        return
    for col_idx, _ in enumerate(ws[1], start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in ws[col_letter])
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)



def apply_positive_negative_formatting(ws) -> None:
    positive_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    negative_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
    for column_index, cell in enumerate(ws[1], start=1):
        header = str(cell.value or "").lower()
        if "ret" not in header and "p&l" not in header and "edge" not in header:
            continue
        column = get_column_letter(column_index)
        cell_range = f"{column}2:{column}{max(ws.max_row, 2)}"
        ws.conditional_formatting.add(cell_range, CellIsRule(operator="greaterThan", formula=["0"], fill=positive_fill))
        ws.conditional_formatting.add(cell_range, CellIsRule(operator="lessThan", formula=["0"], fill=negative_fill))

def replace_sheet_with_df(wb, name: str, df: pd.DataFrame, index: Optional[int] = None) -> None:
    if name in wb.sheetnames:
        old_index = wb.sheetnames.index(name)
        wb.remove(wb[name])
        ws = wb.create_sheet(name, old_index if index is None else index)
    else:
        ws = wb.create_sheet(name, index)
    ws.append(list(df.columns))
    for _, row in df.iterrows():
        ws.append(row.tolist())
    autosize_columns(ws)



def remove_embedded_api_values(workbook) -> None:
    if PARAM_SHEET not in workbook.sheetnames:
        return
    ws = workbook[PARAM_SHEET]
    for row in range(1, ws.max_row + 1):
        label = str(ws.cell(row, 1).value or "").lower()
        if "api key" in label:
            ws.cell(row, 2).value = "Stored as GitHub Actions secret"

def write_output(buynow_df, candidates_df, debug_df, history_df, holdings_df, portfolio_summary_df, template_file, output_file, grouped_date, counters, mode):
    wb = load_workbook(template_file)

    buynow_index = wb.sheetnames.index(BUYNOW_SHEET) if BUYNOW_SHEET in wb.sheetnames else None
    candidates_index = wb.sheetnames.index(CANDIDATES_SHEET) if CANDIDATES_SHEET in wb.sheetnames else None

    replace_sheet_with_df(wb, BUYNOW_SHEET, buynow_df, buynow_index)
    replace_sheet_with_df(wb, CANDIDATES_SHEET, candidates_df, candidates_index)

    run_info = f"momentum_master_v3.py | active_momentum_no_hard_dist_sma50 | real_sma50_sma200 | candidates_sheet_populated | utc={datetime.now(timezone.utc).isoformat()} | grouped_date={grouped_date} | mode={mode}"
    wb[BUYNOW_SHEET]["Z1"] = run_info
    wb[CANDIDATES_SHEET]["Z1"] = run_info

    diag_df = pd.DataFrame(list(counters.items()) + [("mode", mode), ("grouped_date", grouped_date)], columns=["Metric", "Value"])
    replace_sheet_with_df(wb, DIAGNOSTICS_SHEET, diag_df)
    replace_sheet_with_df(wb, DEBUG_SHEET, debug_df)
    replace_sheet_with_df(wb, TRACKING_SHEET, history_df)
    replace_sheet_with_df(wb, HOLDINGS_SHEET, holdings_df)
    replace_sheet_with_df(wb, PORTFOLIO_SUMMARY_SHEET, portfolio_summary_df)
    remove_embedded_api_values(wb)
    wb.save(output_file)


def historical_v3_main() -> None:
    """Frozen pre-Production-v1 runner, retained for historical comparison."""
    setup_logging()
    api_key = os.getenv("POLYGON_API_KEY")
    if not api_key:
        raise EnvironmentError("Missing POLYGON_API_KEY environment variable")
    if not os.path.exists(INPUT_TEMPLATE):
        raise FileNotFoundError(f"Missing input template: {INPUT_TEMPLATE}")

    tickers = read_universe(INPUT_TEMPLATE)
    portfolio_value = read_portfolio_value_sek(INPUT_TEMPLATE)
    client = PolygonClient(api_key)

    latest_date, latest_map, fx = pick_latest_grouped_date(client)
    latest_dt = datetime.fromisoformat(latest_date).date()
    logging.info("Using grouped stock date %s with %d symbols, USDSEK %.4f", latest_date, len(latest_map), fx)

    history_tickers = list(dict.fromkeys(tickers + ["SPY"]))
    recent253 = collect_recent_trading_maps(client, latest_date, history_tickers, latest_map, 253)
    logging.info("Loaded %d recent trading sessions for returns and real SMA calculations.", len(recent253))

    metrics, counters, debug_df = build_metrics(tickers, latest_map, recent253)
    if not metrics:
        raise RuntimeError("Zero valid candidates after active momentum refined processing")

    mode = "normal_top15" if len(metrics) >= TOP_N else f"warning_fewer_than_15_candidates({len(metrics)})"
    buynow = build_buynow_df(metrics, fx, portfolio_value)
    candidates = build_candidates_df(metrics, fx, portfolio_value)
    split_events = update_split_cache(client, latest_date)
    history_df = update_buy_history(buynow, latest_map, latest_date, split_events_by_ticker=split_events)
    portfolio_df = load_portfolio()
    holdings_df = build_holdings_df(portfolio_df, latest_map, fx, candidates, recent253)
    portfolio_summary_df = build_portfolio_summary(holdings_df, latest_date)
    write_output(buynow, candidates, debug_df, history_df, holdings_df, portfolio_summary_df, INPUT_TEMPLATE, OUTPUT_FILE, latest_date, counters, mode)
    logging.info("Wrote %s with %d BuyNow rows and %d candidate rows.", OUTPUT_FILE, len(buynow), len(candidates))


if __name__ == "__main__":
    # Preserve the old command as a safe compatibility alias. It must never
    # execute the retired v3 selector in production.
    from momentum_production import main as production_main

    production_main()
