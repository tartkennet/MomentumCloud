import json
import unittest
import tempfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from momentum_config import PRODUCTION
from update_paper_portfolio import refresh_payload
from momentum_production import (
    _initialize_paper_portfolio, _paper_portfolio, _signals, _write_workbook, fetch_security_master, validate_payload,
)


class ProductionOutputTests(unittest.TestCase):
    def test_security_master_paginates_and_filters(self):
        class Client:
            def __init__(self):
                self.calls = 0

            def _get(self, path, params):
                self.calls += 1
                if self.calls == 1:
                    return {
                        "results": [
                            {"ticker": "GOOD", "name": "Good Inc", "type": "CS", "primary_exchange": "XNAS", "active": True},
                            {"ticker": "FUND", "name": "Fund", "type": "ETF", "primary_exchange": "XNAS", "active": True},
                        ],
                        "next_url": "https://api.polygon.io/v3/reference/tickers?cursor=two",
                    }
                return {"results": [
                    {"ticker": "SPAC", "name": "Example Acquisition Corp", "type": "CS", "primary_exchange": "XNYS", "active": True},
                    {"ticker": "OTHER", "name": "Other Inc", "type": "CS", "primary_exchange": "XNYS", "active": True},
                ]}

        client = Client()
        # The production cache is restored in CI; isolate this pagination unit test
        # so it exercises the client response rather than a prior workflow cache.
        with patch("momentum_production._load_cached_security_master", return_value={}), patch(
            "momentum_production._save_security_master_cache"
        ):
            self.assertEqual(sorted(fetch_security_master(client)), ["GOOD", "OTHER"])
        self.assertEqual(client.calls, 2)

    def test_paper_portfolio_uses_latest_close_and_spy(self):
        with tempfile.TemporaryDirectory() as directory:
            paper_file = Path(directory) / "paper.json"
            paper_file.write_text(json.dumps({
                "status": "ACTIVE",
                "entry_date": "2026-09-15",
                "entry_fx": 10.0,
                "spy_entry_price_usd": 100.0,
                "positions": [
                    {"ticker": "AAA", "quantity": 1, "entry_price_usd": 100.0},
                    {"ticker": "BBB", "quantity": 1, "entry_price_usd": 100.0},
                ],
            }), encoding="utf-8")
            paper = _paper_portfolio(
                {
                    "AAA": {"close": 110.0},
                    "BBB": {"close": 90.0},
                    "SPY": {"close": 105.0},
                },
                "2026-09-16",
                10.0,
                paper_file,
            )
        self.assertEqual(paper["status"], "ACTIVE")
        self.assertEqual(paper["portfolio_return_pct"], 0.0)
        self.assertEqual(paper["spy_return_pct"], 5.0)
        self.assertEqual(paper["vs_spy_pct"], -5.0)

    def test_paper_portfolio_initializes_whole_shares_with_minimum_one(self):
        with tempfile.TemporaryDirectory() as directory:
            paper_file = Path(directory) / "paper.json"
            paper_file.write_text(json.dumps({
                "status": "PENDING_ENTRY",
                "target_allocation_sek": 2_000.0,
                "positions": [{"ticker": "LOW"}, {"ticker": "HIGH"}],
            }), encoding="utf-8")
            _initialize_paper_portfolio(
                {"LOW": {"close": 10.0}, "HIGH": {"close": 300.0}, "SPY": {"close": 100.0}},
                "2026-09-16",
                10.0,
                paper_file,
            )
            paper = _paper_portfolio(
                {"LOW": {"close": 11.0}, "HIGH": {"close": 300.0}, "SPY": {"close": 101.0}},
                "2026-09-17",
                10.0,
                paper_file,
            )
        quantities = {row["ticker"]: row["quantity"] for row in paper["positions"]}
        self.assertEqual(paper["status"], "ACTIVE")
        self.assertEqual(quantities, {"LOW": 20, "HIGH": 1})
        self.assertEqual(paper["entry_value_sek"], 5000.0)

    def test_paper_close_fallback_updates_only_paper_portfolio(self):
        class Client:
            def fetch_daily_open_close(self, ticker, date_str):
                prices = {"AAA": 110.0, "BBB": 90.0, "SPY": 105.0}
                return {"close": prices[ticker]}

            def fetch_fx_close(self, date_str):
                return 10.0

        with tempfile.TemporaryDirectory() as directory:
            paper_file = Path(directory) / "paper.json"
            paper_file.write_text(json.dumps({
                "status": "ACTIVE",
                "entry_date": "2026-09-15",
                "entry_fx": 10.0,
                "spy_entry_price_usd": 100.0,
                "positions": [
                    {"ticker": "AAA", "quantity": 1, "entry_price_usd": 100.0},
                    {"ticker": "BBB", "quantity": 1, "entry_price_usd": 100.0},
                ],
            }), encoding="utf-8")
            payload = {"market_date": "2026-09-18", "buy_now": [{"ticker": "UNCHANGED"}], "system": {"usdsek": 10.0}}
            refreshed = refresh_payload(payload, Client(), paper_file, today=date(2026, 9, 21))
        self.assertEqual(refreshed["market_date"], "2026-09-18")
        self.assertEqual(refreshed["buy_now"], [{"ticker": "UNCHANGED"}])
        self.assertEqual(refreshed["paper_portfolio"]["market_date"], "2026-09-21")
        self.assertEqual(refreshed["paper_portfolio"]["portfolio_return_pct"], 0.0)
        self.assertEqual(refreshed["paper_portfolio"]["spy_return_pct"], 5.0)

    def test_risk_off_liquidates_and_never_buys(self):
        holdings = [{"ticker": "ABC", "stop_triggered": False}]
        target = [{"ticker": "XYZ", "rank": 1, "score": 99, "weight_pct": 100 / 12}]
        self.assertEqual(
            _signals(holdings, target, risk_on=False, rebalance_due=True),
            [{"ticker": "ABC", "action": "SELL", "reason": "RISK_OFF"}],
        )

    def test_monthly_rebalance_and_daily_stop(self):
        holdings = [
            {"ticker": "KEEP", "stop_triggered": False},
            {"ticker": "DROP", "stop_triggered": False},
            {"ticker": "STOP", "stop_triggered": True},
        ]
        target = [
            {"ticker": "KEEP", "rank": 1, "score": 99, "weight_pct": 100 / 12},
            {"ticker": "NEW", "rank": 2, "score": 98, "weight_pct": 100 / 12},
            {"ticker": "STOP", "rank": 3, "score": 97, "weight_pct": 100 / 12},
        ]
        signals = _signals(holdings, target, risk_on=True, rebalance_due=True)
        action = {row["ticker"]: (row["action"], row["reason"]) for row in signals}
        self.assertEqual(action["KEEP"], ("HOLD", "IN_PORTFOLIO"))
        self.assertEqual(action["DROP"], ("SELL", "MONTHLY_REBALANCE"))
        self.assertEqual(action["STOP"], ("SELL", "STOP_LOSS_10_PCT"))
        self.assertEqual(action["NEW"], ("BUY", "MONTHLY_REBALANCE"))

    def test_unfilled_monthly_buy_remains_actionable(self):
        target = [{
            "ticker": "NEW", "rank": 1, "target_rank": 1,
            "score": 99, "weight_pct": 100 / 12,
            "quantity": 3, "price_usd": 20,
        }]
        self.assertEqual(
            _signals([], target, risk_on=True, rebalance_due=False),
            [{
                "ticker": "NEW", "action": "BUY",
                "reason": "PENDING_MONTHLY_REBALANCE", "rank": 1,
                "target_rank": 1, "score": 99,
                "weight_pct": 100 / 12, "quantity": 3,
                "price_usd": 20,
            }],
        )

    def test_stopped_position_is_not_rebought_during_same_month(self):
        target = [{
            "ticker": "STOP", "rank": 1, "target_rank": 1,
            "score": 99, "weight_pct": 100 / 12,
            "quantity": 3, "price_usd": 20,
        }]
        self.assertEqual(
            _signals(
                [], target, risk_on=True, rebalance_due=False,
                blocked_tickers={"STOP"},
            ),
            [],
        )

    def test_payload_guard_rejects_configuration_drift(self):
        target = [
            {
                "ticker": f"T{index}",
                "price_usd": 20.0,
                "adv20_usd": 6_000_000.0,
                "return_1m_pct": 5.0,
                "return_3m_pct": 10.0,
                "mom12_2_pct": 12.0,
                "mom6_2_pct": 8.0,
                "relative_1m_pct": 3.0,
                "relative_3m_pct": 6.0,
                "positive_months": 9,
                "quantity": 1,
                "weight_pct": 100 / 15,
            }
            for index in range(15)
        ]
        payload = {
            "config": {
                "selection_model": PRODUCTION.selection_model,
                "position_count": 15,
                "min_price_usd": 10.0,
                "min_adv20_usd": 5_000_000.0,
                "stop_loss_pct": 10.0,
            },
            "limitation": PRODUCTION.limitation,
            "model_validation": {
                "approved": True,
                "status": "APPROVED",
                "gates": {
                    "at_least_30_periods": True,
                    "positive_net_return": True,
                    "beats_spy": True,
                    "daily_max_drawdown_at_most_20_pct": True,
                    "hit_rate_at_least_50_pct": True,
                },
            },
            "risk": {"status": "RISK ON"},
            "buy_now": [],
            "target_portfolio": target,
            "candidates": [],
        }
        validate_payload(payload)
        payload["config"]["min_adv20_usd"] = 20_000_000.0
        with self.assertRaisesRegex(RuntimeError, "min_adv20_usd"):
            validate_payload(payload)

    def test_unvalidated_model_cannot_publish_buy_now(self):
        payload = {
            "config": {
                "selection_model": PRODUCTION.selection_model,
                "position_count": 15,
                "min_price_usd": 10.0,
                "min_adv20_usd": 5_000_000.0,
                "stop_loss_pct": 10.0,
            },
            "limitation": PRODUCTION.limitation,
            "model_validation": {"approved": False, "status": "UNVALIDATED"},
            "risk": {"status": "RISK ON"},
            "buy_now": [{"ticker": "SHOULD_NOT_APPEAR"}],
            "target_portfolio": [],
            "candidates": [],
        }
        with self.assertRaisesRegex(RuntimeError, "Unvalidated model"):
            validate_payload(payload)

    def test_workbook_contract_also_handles_risk_off_empty_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "template.xlsx"
            output = Path(directory) / "output.xlsx"
            workbook = Workbook()
            workbook.active.title = "Universe"
            workbook.save(template)
            with patch("momentum_production.TEMPLATE_FILE", str(template)), patch(
                "momentum_production.OUTPUT_FILE", str(output)
            ):
                _write_workbook(
                    [], [], [], [], [],
                    {"status": "RISK OFF", "risk_on": False},
                    "2026-09-11",
                )
            generated = load_workbook(output, read_only=True)
            self.assertTrue({"BuyNow", "TargetPortfolio", "Kandidater", "Holdings", "Signals", "PaperPortfolio", "Diagnostics"}.issubset(generated.sheetnames))


if __name__ == "__main__":
    unittest.main()
