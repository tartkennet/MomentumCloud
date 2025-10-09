# momentum_master.py — BuyNow + auto-USDSEK från Polygon + upsert i Parametrar
import os, sys, re, json
import numpy as np
import pandas as pd
from shutil import copyfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import requests

MALL_PATH = "Momentum mall fil.xlsx"
OUT_PATH  = "Momentum.xlsx"
EXCLUDE_TICKERS = {"WOLF"}  # lägg fler tickers här om du vill blocka dem helt

def log(m): print(m, flush=True)
def fail(m): print(f"Error: {m}", file=sys.stderr, flush=True); sys.exit(1)

def to_num_series(s: pd.Series) -> pd.Series:
    if s.dtype.kind in "biufc": return pd.to_numeric(s, errors="coerce")
    x = s.astype(str).str.strip()
    x = x.str.replace("%","",regex=False).str.replace("\u00A0","",regex=False).str.replace(" ","",regex=False)
    x = x.str.replace(",",".",regex=False)
    return pd.to_numeric(x, errors="coerce")

# ---------- 1) Läs Calc ----------
if not os.path.exists(MALL_PATH): fail(f"Mallen saknas: {MALL_PATH}")
try:
    calc = pd.read_excel(MALL_PATH, sheet_name="Calc")
except Exception as e:
    fail(f"Kunde inte läsa fliken 'Calc': {e}")

for must in ["Ticker","Price_0","SMA200"]:
    if must not in calc.columns: fail(f"Saknar kolumn i Calc: '{must}'")

work = pd.DataFrame({"Ticker": calc["Ticker"].astype(str)})
work["Price"]  = to_num_series(calc["Price_0"])
work["SMA200"] = to_num_series(calc["SMA200"])

def safe_ret(p0, p_back):
    p0 = to_num_series(p0); p_back = to_num_series(p_back); return (p0/p_back) - 1.0
if "Price_1m"  in calc.columns: work["Ret_1m"]  = safe_ret(calc["Price_0"], calc["Price_1m"])
if "Price_3m"  in calc.columns: work["Ret_3m"]  = safe_ret(calc["Price_0"], calc["Price_3m"])
if "Price_6m"  in calc.columns: work["Ret_6m"]  = safe_ret(calc["Price_0"], calc["Price_6m"])
if "Price_12m" in calc.columns: work["Ret_12m"] = safe_ret(calc["Price_0"], calc["Price_12m"])

# Rangkolumn
rank_col = None
for c in ["Score_MOMO","ScoreAdj","Score","MomentumScore","Ret_6m","Ret_3m"]:
    if c in calc.columns:
        work[c] = to_num_series(calc[c]); rank_col = c; break
if rank_col is None: rank_col = "Price"

# ---------- 2) Filter + exkludering ----------
excl_set = {t.upper() for t in EXCLUDE_TICKERS}
excl_mask = work["Ticker"].str.upper().isin(excl_set)

flt = (work["Price"] > work["SMA200"].replace({np.nan: np.inf}))
for rc in ["Ret_1m","Ret_3m","Ret_6m","Ret_12m"]:
    if rc in work.columns: flt &= (work[rc].fillna(-1) > 0)

cand = work.loc[flt & (~excl_mask)].dropna(subset=["Ticker"]).copy()
top = cand.sort_values(rank_col, ascending=False).head(15).copy()
need = 15 - len(top)
if need > 0:
    fill = (work["Price"] > work["SMA200"].replace({np.nan: np.inf})) & (~excl_mask)
    for rc in ["Ret_1m","Ret_3m","Ret_6m","Ret_12m"]:
        if rc in work.columns: fill &= (work[rc].fillna(-1) >= 0)
    filler = work.loc[fill & (~work["Ticker"].isin(top["Ticker"]))].sort_values(rank_col, ascending=False).head(need)
    buy_now = pd.concat([top, filler], ignore_index=True)
else:
    buy_now = top.copy()

buy_now = buy_now.drop_duplicates(subset=["Ticker"]).head(15).copy()
buy_now.insert(0,"BuyRank", np.arange(1, len(buy_now)+1))
buy_now["BuyFlag"] = buy_now["BuyRank"] <= 10

# ---------- 3) Läs Parametrar (kapital) ----------
def load_params_dict_from_mall() -> dict:
    try:
        dfp = pd.read_excel(MALL_PATH, sheet_name="Parametrar").iloc[:, :2].copy()
        dfp.columns = ["key","value"]; dfp["key_norm"] = dfp["key"].astype(str).str.strip().str.lower()
        def parse_val(v):
            try: return float(str(v).replace("\u00A0","").replace(" ","").replace(",","."))
            except: return v
        dfp["value"] = dfp["value"].apply(parse_val)
        return dict(zip(dfp["key_norm"], dfp["value"]))
    except Exception:
        return {}

params = load_params_dict_from_mall()
capital = 0.0
for k in ["kapital att handla för","kapital","kapital_att_handla_for"]:
    if k in params:
        try: capital = float(params[k]); break
        except: pass

# ---------- 4) Hämta USDSEK automatiskt via Polygon ----------
def fetch_usdsek_polygon(default=11.0):
    api_key = os.getenv("POLYGON_API_KEY", "").strip()
    if not api_key:
        log("[WARN] POLYGON_API_KEY saknas – använder fallback USDSEK.")
        return default, "fallback"
    try:
        # "prev close" för C:USDSEK
        url = f"https://api.polygon.io/v2/aggs/ticker/C:USDSEK/prev?adjusted=true&apiKey={api_key}"
        r = requests.get(url, timeout=15)
        r.raise_for_status()
        data = r.json()
        # polygon format: results[0].c = close
        close = float(data["results"][0]["c"])
        if close <= 0: raise ValueError("USDSEK close <= 0")
        return close, "polygon"
    except Exception as e:
        log(f"[WARN] Kunde inte hämta USDSEK från Polygon ({e}) – använder fallback {default}.")
        return default, "fallback"

usdsek, usdsek_source = fetch_usdsek_polygon(default=11.0)

# ---------- 5) QtyToBuy (SEK-budget / (USDpris * USDSEK)) ----------
per_ticker_budget_sek = (capital/10.0) if capital and capital>0 else 0.0
buy_now["QtyToBuy"] = 0
if per_ticker_budget_sek > 0 and usdsek > 0:
    mask10 = buy_now["BuyRank"] <= 10
    price_sek = buy_now.loc[mask10,"Price"] * usdsek
    qty = np.floor(per_ticker_budget_sek / price_sek.replace(0,np.nan)).fillna(0).astype(int)
    buy_now.loc[mask10,"QtyToBuy"] = qty

log(f"[INFO] Kapital(SEK)={capital} | USDSEK={usdsek} ({usdsek_source}) | Per-ticker(SEK)={per_ticker_budget_sek} | SumQtyTop10={int(buy_now.loc[buy_now['BuyRank']<=10,'QtyToBuy'].sum())}")

# ---------- 6) Export-DF (format) ----------
export_df = buy_now.copy()
# runda övriga numeriska till 2 dec (ej QtyToBuy/Ret_*)
num_cols = export_df.select_dtypes(include=[np.number]).columns.tolist()
for col in num_cols:
    if col not in ["QtyToBuy","Ret_1m","Ret_3m","Ret_6m","Ret_12m"]:
        export_df[col] = export_df[col].round(2)
export_df = export_df.sort_values("BuyRank").reset_index(drop=True)

# ---------- 7) Kopiera mall → skriv BuyNow EFTER Calc + upsert USDSEK i Parametrar ----------
copyfile(MALL_PATH, OUT_PATH)
wb = load_workbook(OUT_PATH)

# (a) BuyNow efter Calc
if "BuyNow" in wb.sheetnames: wb.remove(wb["BuyNow"])
if "Calc" not in wb.sheetnames: fail("Fliken 'Calc' saknas i utfilen.")
calc_idx = wb.sheetnames.index("Calc")
ws_bn = wb.create_sheet("BuyNow", index=calc_idx+1)

headers = list(export_df.columns)
for j,h in enumerate(headers, start=1): ws_bn.cell(row=1, column=j, value=h)
for i,row in export_df.iterrows():
    for j,h in enumerate(headers, start=1):
        ws_bn.cell(row=i+2, column=j, value=row[h])

# format: procent för returer, heltal för Qty, 0.00 för övrigt
hdr = {h:i+1 for i,h in enumerate(headers)}
def fmt(col, fmt):
    if col not in hdr: return
    c = hdr[col]
    for r in range(2, ws_bn.max_row+1):
        ws_bn.cell(row=r, column=c).number_format = fmt

for c in ["Ret_1m","Ret_3m","Ret_6m","Ret_12m"]:
    fmt(c, "0.00%")
if "QtyToBuy" in hdr: fmt("QtyToBuy","0")
for name in headers:
    if name not in {"Ret_1m","Ret_3m","Ret_6m","Ret_12m","QtyToBuy"}:
        c = hdr[name]
        for r in range(2, ws_bn.max_row+1):
            cell = ws_bn.cell(row=r, column=c)
            if isinstance(cell.value, (int,float)):
                cell.number_format = "0.00"

# (b) Upsert USDSEK i Parametrar
if "Parametrar" in wb.sheetnames:
    ws_p = wb["Parametrar"]
else:
    ws_p = wb.create_sheet("Parametrar")

# hitta om nyckeln redan finns i kol A (key) – vi skriver alltid i två kolumner (A=key, B=value)
def find_row_for_key(sheet, key_name: str):
    for r in range(1, sheet.max_row+1):
        k = sheet.cell(row=r, column=1).value
        if str(k).strip().lower() == key_name.strip().lower():
            return r
    return None

row = find_row_for_key(ws_p, "USDSEK")
if row is None:
    # lägg sist
    row = ws_p.max_row + 1 if ws_p.max_row >= 1 else 1
    ws_p.cell(row=row, column=1, value="USDSEK")
ws_p.cell(row=row, column=2, value=float(usdsek))

# enkel autofit på BuyNow
for col_idx,name in enumerate(headers, start=1):
    maxlen = len(str(name))
    for r in range(2, ws_bn.max_row+1):
        v = ws_bn.cell(row=r, column=col_idx).value
        if v is not None: maxlen = max(maxlen, len(str(v)))
    ws_bn.column_dimensions[get_column_letter(col_idx)].width = min(maxlen+2, 50)

wb.save(OUT_PATH)
log("[OK] Momentum.xlsx uppdaterad — BuyNow efter Calc, USDSEK auto från Polygon och inskriven i Parametrar.")
