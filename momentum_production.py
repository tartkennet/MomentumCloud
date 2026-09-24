#!/usr/bin/env python3
"""Generate all MomentumCloud Production v1 outputs.

This is a decision-support product.  It emits orders/signals but never sends
orders to a broker.  The locked strategy is imported from momentum_strategy;
there are no model or threshold CLI switches.
"""
from __future__ import annotations

import json
import logging
import math
import os
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pandas as pd
from openpyxl import load_workbook

from momentum_config import PRODUCTION
from model_validation import load_model_validation
from momentum_master_v3 import (
    PolygonClient,
    collect_recent_trading_maps,
    get_from_map,
    load_portfolio,
    normalize_ticker,
    pick_latest_grouped_date,
    read_portfolio_value_sek,
    read_universe,
    split_price_factor,
    update_buy_history,
    update_split_cache,
)
from momentum_strategy import (
    Security,
    build_target_portfolio,
    is_eligible_security,
    market_regime,
    rank_universe,
    stop_is_triggered,
)

TEMPLATE_FILE = os.getenv("MOMENTUM_TEMPLATE", "Momentum mall fil.xlsx")
OUTPUT_FILE = os.getenv("MOMENTUM_OUTPUT", "Momentum.xlsx")
STATE_FILE = Path(".momentum_production_state.json")
DASHBOARD_FILE = Path("public/data/dashboard.json")
BUYNOW_FILE = Path("public/data/buynow.json")
SIGNALS_FILE = Path("Momentum_signals.csv")
PERFORMANCE_FILE = Path("Momentum_performance.csv")
PAPER_PORTFOLIO_FILE = Path("Momentum_paper_portfolio.json")
STATE_SCHEMA_VERSION = 3
SECURITY_MASTER_CACHE_FILE = Path(".momentum_security_master.json")
SECURITY_MASTER_CACHE_VERSION = 1
SECURITY_MASTER_CACHE_MAX_AGE_DAYS = 7


def _write_json(path: Path, payload: dict | list) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def _safe_number(value):
    if value is None or pd.isna(value):
        return None
    return float(value)


def _load_cached_security_master(
    cache_file: Path = SECURITY_MASTER_CACHE_FILE,
) -> dict[str, Security]:
    """Use a recent local master to stay within the free API plan's call limit."""
    try:
        payload = json.loads(cache_file.read_text(encoding="utf-8"))
        if payload.get("version") != SECURITY_MASTER_CACHE_VERSION:
            return {}
        cached_at = datetime.fromisoformat(
            str(payload["cached_at"]).replace("Z", "+00:00")
        )
        if cached_at.tzinfo is None:
            cached_at = cached_at.replace(tzinfo=timezone.utc)
        if datetime.now(timezone.utc) - cached_at > timedelta(
            days=SECURITY_MASTER_CACHE_MAX_AGE_DAYS
        ):
            return {}
        securities = {}
        for row in payload.get("securities", []):
            security = Security(
                ticker=normalize_ticker(row.get("ticker", "")),
                name=str(row.get("name", "")),
                security_type=str(row.get("security_type", "")),
                primary_exchange=str(row.get("primary_exchange", "")),
                active=bool(row.get("active", True)),
            )
            if security.ticker and is_eligible_security(security):
                securities[security.ticker] = security
        # A production US security master has thousands of names.  A tiny cache
        # is a test fixture or a partial write, never a valid live universe.
        if len(securities) < 100:
            logging.warning(
                "Discarding unsafe security-master cache with only %d securities",
                len(securities),
            )
            return {}
        return securities
    except (AttributeError, KeyError, OSError, TypeError, ValueError, json.JSONDecodeError):
        return {}


def _save_security_master_cache(
    securities: dict[str, Security],
    cache_file: Path = SECURITY_MASTER_CACHE_FILE,
) -> None:
    payload = {
        "version": SECURITY_MASTER_CACHE_VERSION,
        "cached_at": datetime.now(timezone.utc).isoformat(),
        "securities": [
            {
                "ticker": item.ticker,
                "name": item.name,
                "security_type": item.security_type,
                "primary_exchange": item.primary_exchange,
                "active": item.active,
            }
            for item in securities.values()
        ],
    }
    cache_file.write_text(
        json.dumps(payload, separators=(",", ":")),
        encoding="utf-8",
    )


def fetch_security_master(client: PolygonClient) -> dict[str, Security]:
    """Fetch the active US common-stock master once and cache it for seven days."""
    cached = _load_cached_security_master()
    if cached:
        logging.info("Using cached security master with %d eligible securities", len(cached))
        return cached

    path = "/v3/reference/tickers"
    params = {
        "market": "stocks",
        "active": "true",
        "order": "asc",
        "limit": 1000,
        "sort": "ticker",
    }
    securities: dict[str, Security] = {}
    while True:
        data = client._get(path, params)
        for row in data.get("results", []) or []:
            ticker = normalize_ticker(row.get("ticker", ""))
            if not ticker:
                continue
            security = Security(
                ticker=ticker,
                name=str(row.get("name", "")),
                security_type=str(row.get("type", "")),
                primary_exchange=str(row.get("primary_exchange", "")),
                active=row.get("active") is not False,
            )
            if is_eligible_security(security):
                securities[ticker] = security
        next_url = data.get("next_url")
        if not next_url:
            break
        parsed = urlparse(str(next_url))
        path = parsed.path
        params = {
            key: values[-1]
            for key, values in parse_qs(parsed.query).items()
            if key.lower() != "apikey" and values
        }
    if not securities:
        raise RuntimeError("Polygon returned no eligible active US common stocks")
    _save_security_master_cache(securities)
    return securities

def _ranked_rows(
    ranked,
    portfolio_value_sek: float,
    fx: float,
    target: list[dict],
) -> list[dict]:
    position_value = portfolio_value_sek / PRODUCTION.position_count
    target_by_ticker = {row["ticker"]: row for row in target}
    rows = []
    for rank, metric in enumerate(ranked[: PRODUCTION.candidate_count], 1):
        price_sek = metric.price_usd * fx
        affordable_quantity = math.floor(position_value / price_sek) if price_sek > 0 else 0
        target_row = target_by_ticker.get(metric.ticker)
        rows.append(
            {
                "rank": rank,
                "target_rank": target_row.get("target_rank") if target_row else None,
                "ticker": metric.ticker,
                "score": round(metric.score, 4),
                "price_usd": round(metric.price_usd, 4),
                "adv20_usd": round(metric.adv20_usd, 2),
                "weight_pct": round(float(target_row["weight_pct"]), 4)
                if target_row else 0.0,
                "quantity": int(target_row["quantity"]) if target_row else 0,
                "affordable_quantity": affordable_quantity,
                "is_target": target_row is not None,
                "selection_status": (
                    "TARGET" if target_row else
                    "UNAFFORDABLE" if affordable_quantity < 1 else
                    "RESERVE"
                ),
                "return_1m_pct": round(metric.return_1m_pct, 4),
                "return_3m_pct": round(metric.return_3m_pct, 4),
                "return_6m_pct": round(metric.return_6m_pct, 4),
                "mom12_2_pct": round(metric.mom12_2_pct, 4),
                "mom6_2_pct": round(metric.mom6_2_pct, 4),
                "relative_1m_pct": round(metric.relative_1m_pct, 4),
                "relative_3m_pct": round(metric.relative_3m_pct, 4),
                "volatility_pct": round(metric.volatility_pct, 4),
                "extension_sma50_pct": round(metric.extension_sma50_pct, 4),
            }
        )
    return rows


def _holdings(
    portfolio: pd.DataFrame,
    latest_map: dict,
    latest_date: str,
    fx: float,
    split_events: dict,
    target_tickers: set[str],
) -> list[dict]:
    rows = []
    for holding in portfolio.to_dict(orient="records"):
        ticker = normalize_ticker(holding["Ticker"])
        market = get_from_map(latest_map, ticker)
        entry = float(holding["EntryPrice_USD"])
        adjusted_entry = entry * split_price_factor(
            split_events.get(ticker, []), str(holding["EntryDate"]), latest_date
        )
        price = float(market["close"]) if market else None
        return_pct = (price / adjusted_entry - 1.0) * 100.0 if price else None
        rows.append(
            {
                "ticker": ticker,
                "quantity": float(holding["Qty"]),
                "entry_date": str(holding["EntryDate"]),
                "entry_price_usd": entry,
                "adjusted_entry_price_usd": adjusted_entry,
                "current_price_usd": price,
                "return_pct": return_pct,
                "pnl_sek": (
                    float(holding["Qty"]) * price * fx
                    - float(holding["Qty"]) * entry * float(holding["EntryFX"])
                    - float(holding.get("EntryFees_SEK", 0.0) or 0.0)
                )
                if price
                else None,
                "in_target": ticker in target_tickers,
                "stop_triggered": stop_is_triggered(price, adjusted_entry) if price else False,
            }
        )
    return rows


def _signals(
    holdings: list[dict],
    target: list[dict],
    risk_on: bool,
    rebalance_due: bool,
    blocked_tickers: set[str] | None = None,
) -> list[dict]:
    blocked_tickers = blocked_tickers or set()
    held = {row["ticker"]: row for row in holdings}
    target_by_ticker = {row["ticker"]: row for row in target}
    signals = []

    for ticker, row in held.items():
        if not risk_on:
            action, reason = "SELL", "RISK_OFF"
        elif row["stop_triggered"]:
            action, reason = "SELL", "STOP_LOSS_10_PCT"
        elif rebalance_due and ticker not in target_by_ticker:
            action, reason = "SELL", "MONTHLY_REBALANCE"
        else:
            action, reason = "HOLD", "IN_PORTFOLIO"
        signals.append({"ticker": ticker, "action": action, "reason": reason})

    if risk_on:
        for ticker, row in target_by_ticker.items():
            if ticker not in held and ticker not in blocked_tickers:
                signals.append(
                    {
                        "ticker": ticker,
                        "action": "BUY",
                        "reason": (
                            "MONTHLY_REBALANCE" if rebalance_due
                            else "PENDING_MONTHLY_REBALANCE"
                        ),
                        "rank": row["rank"],
                        "target_rank": row.get("target_rank"),
                        "score": row["score"],
                        "weight_pct": row["weight_pct"],
                        "quantity": row.get("quantity"),
                        "price_usd": row.get("price_usd"),
                    }
                )
    return sorted(signals, key=lambda row: ({"SELL": 0, "BUY": 1, "HOLD": 2}[row["action"]], row["ticker"]))


def _performance(recent_maps, holdings: list[dict]) -> dict:
    spy = []
    for date, day_map in recent_maps:
        row = get_from_map(day_map, "SPY")
        if row:
            spy.append((date, float(row["close"])))
    spy_1m = (spy[0][1] / spy[21][1] - 1.0) * 100.0 if len(spy) > 21 else None
    spy_12m = (spy[0][1] / spy[252][1] - 1.0) * 100.0 if len(spy) > 252 else None

    actual_returns = [row["return_pct"] for row in holdings if row["return_pct"] is not None]
    current_return = sum(actual_returns) / len(actual_returns) if actual_returns else None
    return {
        "portfolio_return_pct": current_return,
        "spy_1m_pct": spy_1m,
        "spy_12m_pct": spy_12m,
        "cagr_pct": None,
        "sharpe": None,
        "sortino": None,
        "max_drawdown_pct": None,
        "status": "Awaiting sufficient live portfolio history"
        if not PERFORMANCE_FILE.exists()
        else "Available in performance history",
    }


def _initialize_paper_portfolio(
    latest_map: dict,
    latest_date: str,
    fx: float,
    paper_file: Path = PAPER_PORTFOLIO_FILE,
) -> dict:
    """Convert the user's requested paper allocation into whole-share holdings once."""
    try:
        request = json.loads(paper_file.read_text(encoding="utf-8"))
        if request.get("status") == "ACTIVE":
            return request
        allocation_sek = float(request["target_allocation_sek"])
        requested = request["positions"]
        if allocation_sek <= 0 or not requested:
            raise ValueError("paper portfolio request is incomplete")
    except (KeyError, OSError, TypeError, ValueError, json.JSONDecodeError) as exc:
        logging.warning("Paper portfolio was not initialized: %s", exc)
        return {}

    spy_market = get_from_map(latest_map, "SPY")
    if not spy_market or float(spy_market.get("close") or 0.0) <= 0.0:
        logging.warning("Paper portfolio is waiting for a verified SPY close")
        return request

    positions = []
    for item in requested:
        ticker = normalize_ticker(item.get("ticker", ""))
        market = get_from_map(latest_map, ticker)
        price = float(market["close"]) if market else 0.0
        if not ticker or price <= 0.0:
            logging.warning("Paper portfolio is waiting for a verified close for %s", ticker or "unknown")
            return request
        quantity = max(1, math.floor(allocation_sek / (price * fx)))
        positions.append(
            {
                "ticker": ticker,
                "quantity": int(quantity),
                "entry_price_usd": round(price, 4),
                "entry_value_sek": round(quantity * price * fx, 2),
            }
        )

    portfolio = {
        "schema_version": 2,
        "status": "ACTIVE",
        "entry_date": latest_date,
        "entry_basis": "Verified closing prices from the first production snapshot after the paper request",
        "allocation_rule": "Approximately 2,000 SEK per position; at least one whole share",
        "target_allocation_sek": round(allocation_sek, 2),
        "entry_fx": round(float(fx), 6),
        "spy_entry_price_usd": round(float(spy_market["close"]), 4),
        "positions": positions,
    }
    _write_json(paper_file, portfolio)
    return portfolio


def _paper_portfolio(
    latest_map: dict,
    latest_date: str,
    fx: float,
    paper_file: Path = PAPER_PORTFOLIO_FILE,
) -> dict:
    """Mark the user's whole-share paper portfolio against its verified entry close."""
    try:
        seed = json.loads(paper_file.read_text(encoding="utf-8"))
        entries = seed["positions"]
        if not entries:
            raise ValueError("paper portfolio has no positions")
    except (KeyError, OSError, TypeError, ValueError, json.JSONDecodeError) as exc:
        return {
            "status": "NOT_STARTED",
            "market_date": latest_date,
            "reason": f"Paper portfolio unavailable: {exc}",
            "positions": [],
        }

    if seed.get("status") != "ACTIVE":
        return {
            "status": "PENDING_ENTRY",
            "market_date": latest_date,
            "position_count": len(entries),
            "target_allocation_sek": seed.get("target_allocation_sek"),
            "positions": [{"ticker": normalize_ticker(item.get("ticker", ""))} for item in entries],
            "note": "Awaiting the next verified production close to set whole-share paper entries.",
        }

    try:
        entry_date = str(seed["entry_date"])
        entry_fx = float(seed["entry_fx"])
        spy_entry = float(seed["spy_entry_price_usd"])
    except (KeyError, TypeError, ValueError) as exc:
        return {
            "status": "NOT_STARTED",
            "market_date": latest_date,
            "reason": f"Paper portfolio entry is incomplete: {exc}",
            "positions": [],
        }

    rows = []
    entry_total_sek = 0.0
    current_total_sek = 0.0
    for item in entries:
        ticker = normalize_ticker(item.get("ticker", ""))
        quantity = int(item["quantity"])
        entry_price = float(item["entry_price_usd"])
        entry_value = float(item.get("entry_value_sek") or quantity * entry_price * entry_fx)
        market = get_from_map(latest_map, ticker)
        current_price = float(market["close"]) if market else None
        current_value = quantity * current_price * fx if current_price is not None else None
        price_return = ((current_price / entry_price) - 1.0) * 100.0 if current_price is not None else None
        return_pct = ((current_value / entry_value) - 1.0) * 100.0 if current_value is not None else None
        entry_total_sek += entry_value
        if current_value is not None:
            current_total_sek += current_value
        rows.append(
            {
                "ticker": ticker,
                "quantity": quantity,
                "entry_price_usd": round(entry_price, 4),
                "current_price_usd": round(current_price, 4) if current_price is not None else None,
                "entry_value_sek": round(entry_value, 2),
                "current_value_sek": round(current_value, 2) if current_value is not None else None,
                "pnl_sek": round(current_value - entry_value, 2) if current_value is not None else None,
                "price_return_pct": round(price_return, 4) if price_return is not None else None,
                "return_pct": round(return_pct, 4) if return_pct is not None else None,
                "stop_triggered": stop_is_triggered(current_price, entry_price) if current_price is not None else False,
            }
        )

    spy_market = get_from_map(latest_map, "SPY")
    spy_current = float(spy_market["close"]) if spy_market else None
    spy_return = ((spy_current / spy_entry) - 1.0) * 100.0 if spy_current is not None else None
    complete = all(row["current_value_sek"] is not None for row in rows)
    portfolio_return = (
        ((current_total_sek / entry_total_sek) - 1.0) * 100.0
        if complete and entry_total_sek > 0.0
        else None
    )
    return {
        "status": "ACTIVE",
        "entry_date": entry_date,
        "entry_basis": seed.get("entry_basis", "verified closing prices in USD"),
        "allocation_rule": seed.get("allocation_rule"),
        "target_allocation_sek": seed.get("target_allocation_sek"),
        "entry_fx": round(entry_fx, 6),
        "market_date": latest_date,
        "position_count": len(rows),
        "entry_value_sek": round(entry_total_sek, 2),
        "current_value_sek": round(current_total_sek, 2) if complete else None,
        "pnl_sek": round(current_total_sek - entry_total_sek, 2) if complete else None,
        "portfolio_return_pct": round(portfolio_return, 4) if portfolio_return is not None else None,
        "spy_entry_price_usd": round(spy_entry, 4),
        "spy_current_price_usd": round(spy_current, 4) if spy_current is not None else None,
        "spy_return_pct": round(spy_return, 4) if spy_return is not None else None,
        "vs_spy_pct": round(portfolio_return - spy_return, 4)
        if portfolio_return is not None and spy_return is not None else None,
        "stopped_tickers": sorted(row["ticker"] for row in rows if row["stop_triggered"]),
        "positions": rows,
        "note": "Paper tracking only. No orders are sent. Returns include USD/SEK movement and exclude fees and taxes.",
    }

def _history() -> list[dict]:
    path = Path("Momentum_buy_history.csv")
    if not path.exists():
        return []
    frame = pd.read_csv(path).tail(100).copy()
    return json.loads(frame.to_json(orient="records", date_format="iso"))


def _excel_frame(rows: list[dict]) -> pd.DataFrame:
    if not rows:
        return pd.DataFrame(
            columns=[
                "Rank", "Ticker", "TotalScore", "Price_USD",
                "AvgDollarVolume20", "Weight_pct", "Qty",
                "Return_1m_pct", "Return_3m_pct", "Return_6m_pct",
                "Relative_1m_pct", "Relative_3m_pct",
                "Volatility_pct", "Extension_SMA50_pct", "buy_signal",
            ]
        )
    return pd.DataFrame(rows).rename(
        columns={
            "rank": "Rank",
            "ticker": "Ticker",
            "score": "TotalScore",
            "price_usd": "Price_USD",
            "adv20_usd": "AvgDollarVolume20",
            "weight_pct": "Weight_pct",
            "quantity": "Qty",
            "return_1m_pct": "Return_1m_pct",
            "return_3m_pct": "Return_3m_pct",
            "return_6m_pct": "Return_6m_pct",
            "relative_1m_pct": "Relative_1m_pct",
            "relative_3m_pct": "Relative_3m_pct",
            "volatility_pct": "Volatility_pct",
            "extension_sma50_pct": "Extension_SMA50_pct",
        }
    )


def _replace_sheet(workbook, name: str, frame: pd.DataFrame) -> None:
    index = workbook.sheetnames.index(name) if name in workbook.sheetnames else len(workbook.sheetnames)
    if name in workbook.sheetnames:
        workbook.remove(workbook[name])
    worksheet = workbook.create_sheet(name, index)
    worksheet.append(list(frame.columns) or ["Status"])
    for row in frame.itertuples(index=False, name=None):
        worksheet.append(list(row))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def _write_workbook(
    buy_now: list[dict],
    target: list[dict],
    candidates: list[dict],
    holdings: list[dict],
    signals: list[dict],
    risk: dict,
    latest_date: str,
    paper_portfolio: dict | None = None,
) -> None:
    workbook = load_workbook(TEMPLATE_FILE)
    _replace_sheet(workbook, "BuyNow", _excel_frame(buy_now))
    _replace_sheet(workbook, "TargetPortfolio", _excel_frame(target))
    _replace_sheet(workbook, "Kandidater", _excel_frame(candidates))
    _replace_sheet(workbook, "Holdings", pd.DataFrame(holdings))
    _replace_sheet(workbook, "Signals", pd.DataFrame(signals))
    _replace_sheet(
        workbook,
        "PaperPortfolio",
        pd.DataFrame((paper_portfolio or {}).get("positions", [])),
    )
    diagnostics = pd.DataFrame(
        [
            ("Product", PRODUCTION.product_name),
            ("Selection model", PRODUCTION.selection_model),
            ("Freshness rule", "1m and 3m positive and above SPY"),
            ("Market date", latest_date),
            ("Regime", risk["status"]),
            ("Positions", PRODUCTION.position_count),
            ("Minimum price USD", PRODUCTION.min_price_usd),
            ("Minimum ADV20 USD", PRODUCTION.min_adv20_usd),
            ("Stop loss pct", PRODUCTION.stop_loss_pct),
            ("Limitation", PRODUCTION.limitation),
        ],
        columns=["Metric", "Value"],
    )
    _replace_sheet(workbook, "Diagnostics", diagnostics)
    workbook.save(OUTPUT_FILE)


def validate_payload(payload: dict) -> None:
    config = payload["config"]
    expected = {
        "selection_model": "v4-balanced",
        "position_count": 15,
        "min_price_usd": 10.0,
        "min_adv20_usd": 5_000_000.0,
        "stop_loss_pct": 10.0,
    }
    for key, value in expected.items():
        if config.get(key) != value:
            raise RuntimeError(f"Production config mismatch: {key}={config.get(key)!r}")
    if payload["limitation"] != "Current limitation: survivorship bias":
        raise RuntimeError("Required survivorship-bias limitation is missing")

    validation = payload.get("model_validation") or {}
    approved = validation.get("approved") is True
    target = payload.get("target_portfolio", [])
    buys = payload.get("buy_now", [])
    if not approved:
        if target or buys:
            raise RuntimeError("Unvalidated model must not publish target or Buy Now rows")
        return

    required_gates = {
        "at_least_30_periods",
        "positive_net_return",
        "beats_spy",
        "daily_max_drawdown_at_most_20_pct",
        "hit_rate_at_least_50_pct",
    }
    if validation.get("status") != "APPROVED" or not all(
        validation.get("gates", {}).get(gate) is True for gate in required_gates
    ):
        raise RuntimeError("Approved model is missing a complete verified validation decision")

    risk_on = payload["risk"]["status"] == "RISK ON"
    if not risk_on and (buys or target):
        raise RuntimeError("RISK OFF output must not contain targets or Buy Now orders")
    if risk_on:
        if len(target) != PRODUCTION.position_count:
            raise RuntimeError("RISK ON requires exactly 15 executable target positions")
        if len({row["ticker"] for row in target}) != PRODUCTION.position_count:
            raise RuntimeError("Target portfolio contains duplicate tickers")
        if any(int(row.get("quantity") or 0) < 1 for row in target):
            raise RuntimeError("Every target position must contain at least one whole share")
        if any(float(row["price_usd"]) < PRODUCTION.min_price_usd for row in target):
            raise RuntimeError("Target portfolio contains a stock below the price floor")
        if any(float(row["adv20_usd"]) < PRODUCTION.min_adv20_usd for row in target):
            raise RuntimeError("Target portfolio contains a stock below the ADV20 floor")
        if any(float(row["return_1m_pct"]) <= 0.0 for row in target):
            raise RuntimeError("Target portfolio contains a stock with non-positive 1-month return")
        if any(float(row["return_3m_pct"]) <= 0.0 for row in target):
            raise RuntimeError("Target portfolio contains a stock with non-positive 3-month return")
        if any(float(row["relative_1m_pct"]) <= 0.0 for row in target):
            raise RuntimeError("Target portfolio contains a stock lagging SPY over 1 month")
        if any(float(row["relative_3m_pct"]) <= 0.0 for row in target):
            raise RuntimeError("Target portfolio contains a stock lagging SPY over 3 months")
        if abs(sum(float(row["weight_pct"]) for row in target) - 100.0) > 0.01:
            raise RuntimeError("Target weights do not sum to 100 percent")
    target_tickers = {row["ticker"] for row in target}
    if any(row.get("ticker") not in target_tickers for row in buys):
        raise RuntimeError("Buy Now contains a ticker outside the locked target portfolio")
    if any(row.get("buy_signal") != "BUY" for row in buys):
        raise RuntimeError("Buy Now may contain actionable BUY orders only")
    if len(buys) > 15 or len(payload["candidates"]) > 50:
        raise RuntimeError("Production portfolio or candidate limit exceeded")

def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
    api_key = os.getenv("POLYGON_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("POLYGON_API_KEY is required")
    PRODUCTION.validate()

    client = PolygonClient(api_key)
    template_universe = read_universe(TEMPLATE_FILE)
    security_master = fetch_security_master(client)
    eligible = [ticker for ticker in template_universe if normalize_ticker(ticker) in security_master]
    if len(eligible) < PRODUCTION.position_count:
        raise RuntimeError(
            f"Only {len(eligible)} eligible US common stocks remain after security filtering"
        )

    latest_date, latest_map, fx = pick_latest_grouped_date(client)
    _initialize_paper_portfolio(latest_map, latest_date, fx)
    history_tickers = list(dict.fromkeys([*eligible, "SPY"]))
    recent_maps = collect_recent_trading_maps(
        client, latest_date, history_tickers, latest_map, needed=253
    )
    risk = market_regime(recent_maps)
    previous_state = {}
    if STATE_FILE.exists():
        previous_state = json.loads(STATE_FILE.read_text(encoding="utf-8"))
    selection_month = latest_date[:7]
    # Re-enter promptly after a risk-off interval; otherwise rotate once per
    # month. This is a regime transition, not an extra ranking cycle.
    state_is_current = (
        previous_state.get("schema_version") == STATE_SCHEMA_VERSION
        and isinstance(previous_state.get("target_portfolio"), list)
        and len(previous_state.get("target_portfolio", [])) == PRODUCTION.position_count
    )
    rebalance_due = (
        not state_is_current
        or previous_state.get("selection_month") != selection_month
        or (
            risk["risk_on"]
            and previous_state.get("risk_status") not in {None, "RISK ON"}
        )
    )

    ranked = rank_universe(eligible, recent_maps)
    portfolio_value_sek = read_portfolio_value_sek(TEMPLATE_FILE)
    model_validation = load_model_validation()
    if risk["risk_on"] and model_validation["approved"] and rebalance_due:
        target = build_target_portfolio(
            ranked,
            True,
            portfolio_value_sek=portfolio_value_sek,
            usdsek=fx,
        )
    elif risk["risk_on"] and model_validation["approved"]:
        target = [dict(row) for row in previous_state["target_portfolio"]]
    else:
        # Candidates remain visible for research, but an unvalidated model must
        # never emit a target portfolio or actionable BUY signal.
        target = []
    target_tickers = {row["ticker"] for row in target}
    ranked_rows = _ranked_rows(ranked, portfolio_value_sek, fx, target)

    split_events = update_split_cache(client, latest_date)
    portfolio = load_portfolio()
    holdings = _holdings(portfolio, latest_map, latest_date, fx, split_events, target_tickers)
    blocked_tickers = set()
    if not rebalance_due:
        blocked_tickers.update(previous_state.get("blocked_tickers", []))
    blocked_tickers.update(
        row["ticker"] for row in holdings if row.get("stop_triggered")
    )
    signals = _signals(
        holdings,
        target,
        risk["risk_on"],
        rebalance_due,
        blocked_tickers,
    )
    buy_tickers = {
        row["ticker"] for row in signals if row.get("action") == "BUY"
    }
    buy_now = [
        {**dict(row), "buy_signal": "BUY"}
        for row in target
        if row["ticker"] in buy_tickers
    ]

    run_url = None
    if os.getenv("GITHUB_RUN_ID") and os.getenv("GITHUB_REPOSITORY"):
        run_url = f"https://github.com/{os.environ['GITHUB_REPOSITORY']}/actions/runs/{os.environ['GITHUB_RUN_ID']}"
    payload = {
        "product": PRODUCTION.product_name,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "market_date": latest_date,
        "selection_month": selection_month,
        "rebalance_due": rebalance_due,
        "config": PRODUCTION.to_dict(),
        "limitation": PRODUCTION.limitation,
        "risk": {**risk, "exposure_pct": 100.0 if risk["risk_on"] and model_validation["approved"] else 0.0},
        "model_validation": model_validation,
        "buy_now": buy_now,
        "target_portfolio": target,
        "candidates": ranked_rows,
        "holdings": holdings,
        "signals": signals,
        "performance": _performance(recent_maps, holdings),
        "paper_portfolio": _paper_portfolio(latest_map, latest_date, fx),
        "history": _history(),
        "system": {
            "github_action_status": "SUCCESS",
            "github_action_url": run_url,
            "polygon_status": "CONNECTED",
            "universe_input_count": len(template_universe),
            "eligible_security_count": len(eligible),
            "ranked_candidate_count": len(ranked),
            "usdsek": round(float(fx), 6),
        },
    }
    validate_payload(payload)

    _write_json(DASHBOARD_FILE, payload)
    _write_json(BUYNOW_FILE, {"updated": payload["generated_at"], "data": buy_now})
    pd.DataFrame(signals).to_csv(SIGNALS_FILE, index=False)
    _write_workbook(
        buy_now, target, ranked_rows, holdings, signals, risk, latest_date,
        paper_portfolio=payload["paper_portfolio"],
    )
    # Entry history records the monthly recommendation once.  Pending orders
    # remain visible in Buy Now, but must not create a new fictitious entry on
    # every daily refresh.
    history_input = buy_now if rebalance_due else []
    history_frame = update_buy_history(
        _excel_frame(history_input),
        latest_map,
        latest_date,
        split_events_by_ticker=split_events,
    )
    payload["history"] = json.loads(
        history_frame.tail(100).to_json(orient="records", date_format="iso")
    )
    _write_json(DASHBOARD_FILE, payload)
    _write_json(
        STATE_FILE,
        {
            "schema_version": STATE_SCHEMA_VERSION,
            "selection_month": selection_month,
            "market_date": latest_date,
            "risk_status": risk["status"],
            "target_tickers": sorted(target_tickers),
            "target_portfolio": target,
            "blocked_tickers": sorted(blocked_tickers),
        },
    )
    logging.info(
        "Production output ready: %s, %d Buy Now, %d candidates, %s",
        latest_date,
        len(buy_now),
        len(ranked_rows),
        risk["status"],
    )


if __name__ == "__main__":
    main()
